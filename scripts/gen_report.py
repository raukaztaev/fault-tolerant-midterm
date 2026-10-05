#!/usr/bin/env python3
"""Generates the final midterm technical report as a plain docx file."""
import json
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
EX = os.path.join(ROOT, "experiments")
DOCS = os.path.join(ROOT, "docs")

with open(os.path.join(EX, "analysis.json")) as f:
    A = json.load(f)

BL = A["baseline"]
FT = A["fault_tolerant"]
RBD = A["rbd"]
CMP = A["comparison"]

doc = Document()

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

# make headings plain black
for level, size in [(1, 14), (2, 12)]:
    s = doc.styles[f"Heading {level}"]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def p(text):
    doc.add_paragraph(text)


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def lines(items):
    for it in items:
        par = doc.add_paragraph()
        r = par.add_run(it)
        r.font.size = Pt(10)
        par.paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- title page
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("MIDTERM PROJECT\nFault-Tolerant University Information System")
r.bold = True
r.font.size = Pt(18)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Course: Fault Tolerance and Dependable Computing\n"
              "Kaztayev Rauan Daurenuly, ID 255387, group CSE-2507M\n"
              "05.10.2026")
r.font.size = Pt(12)

doc.add_paragraph()

# ---------------------------------------------------------------- 1
h1("1. Introduction and Project Motivation")
p("University digital services must keep working during registration weeks, tuition payment "
  "deadlines and exams. These are also the periods with the highest load, so failures happen "
  "exactly when the service is needed most. The goal of this project is to design, implement "
  "and measure a small university information system that survives realistic failures: crashed "
  "instances, database outages, slow calls, interrupted payments and heavy load.")
p("I built two configurations of the same system. The baseline version has no fault tolerance "
  "mechanisms. The fault-tolerant version has service replication, database replication with "
  "failover, retries, timeouts, a circuit breaker, idempotent payments, checkpointing with "
  "journal recovery and graceful degradation. Both configurations went through the same six "
  "failure scenarios, and every number in this report comes from these experiments.")
p("All data in the project are synthetic educational data and do not describe real AITU "
  "infrastructure.")

# ---------------------------------------------------------------- 2
h1("2. System Requirements and Assumptions")
h2("2.1 Functional scope")
p("The platform has four services behind one gateway on port 8080:")
bullets([
    "Student Service: registration and lookup of student records.",
    "Payment Service: tuition payments with a transaction journal and checkpoints.",
    "Academic Records Service: grades and transcript generation with GPA.",
    "Timetable Service: weekly schedule with a registration check.",
])
h2("2.2 Requirements")
lines([
    "- survive one crashed service instance without user-visible errors (verified in E1)",
    "- survive primary database failure with at least 99% successful payments (E2)",
    "- survive a 2.5 second downstream delay without failed requests (E3)",
    "- survive a paused node and keep serving (E4)",
    "- never lose or duplicate an interrupted payment (E5)",
    "- keep at least 98% success under 24 concurrent workers (E6)",
    "- measured request availability in fault-tolerant mode: at least 99%",
])
h2("2.3 Assumptions")
bullets([
    "The system runs on one machine in Docker Desktop, so a node is a set of containers and node "
    "failure is simulated by pausing a container.",
    "Database replication is implemented at application level (dual write to primary and standby) "
    "instead of native PostgreSQL streaming replication, so the failover logic is visible in Java code.",
    "Two instances per service means the system tolerates one failed instance per service.",
    "A client that retries a payment uses the same idempotency key, as real payment APIs require.",
])

# ---------------------------------------------------------------- 3
h1("3. Dependability and Fault Model")
p("A fault is the cause (a killed container, a stopped database, an injected delay). An error is "
  "the internal state it produces (an unreachable JDBC connection, a timed out HTTP call). A "
  "failure is what the user sees (a 503 response, a lost payment). The fault tolerance mechanisms "
  "in this project break the chain between fault and failure: faults still happen, but they stop "
  "producing user-visible failures.")
lines([
    "1. Application crash (docker kill of one instance). Without FT all requests to it fail. "
    "With FT the gateway fails over to the second replica.",
    "2. Primary database stopped. Without FT every write fails. With FT the standby is promoted "
    "and payments continue after a short timeout window.",
    "3. Service delay of 2.5 seconds. With FT the gateway timeout plus retry reaches the healthy "
    "replica.",
    "4. Node failure (paused container). With FT round robin skips the paused replica.",
    "5. Interrupted payment (instance killed mid transaction). With FT the journal and idempotency "
    "key guarantee a clean replay without a double charge.",
    "6. High load, 24 concurrent workers. With FT the load is spread over two replicas.",
])
p("The experiments cover the crash fault model (E1, E5), the omission and timing fault model "
  "(E2, E3) and the load fault model (E6). Byzantine faults are out of scope.")

# ---------------------------------------------------------------- 4
h1("4. Reliability and Failure Analysis")
h2("4.1 Measured metrics")
p("Each configuration was observed during the six-experiment campaign (145 seconds of injected "
  "failure windows, more than 3400 baseline and 7000 fault-tolerant requests). Downtime was "
  "estimated from the timestamps of failed requests in the request logs.")
lines([
    f"Total requests: baseline {BL['total_requests']}, fault-tolerant {FT['total_requests']}",
    f"Successful requests: baseline {BL['successful']}, fault-tolerant {FT['successful']}",
    f"Failed requests: baseline {BL['failed']}, fault-tolerant {FT['failed']}",
    f"Request availability (ok/total): baseline {BL['request_availability_percent']}%, "
    f"fault-tolerant {FT['request_availability_percent']}%",
    f"Total downtime: baseline {BL['total_downtime_s']} s, fault-tolerant {FT['total_downtime_s']} s",
    f"Time availability (uptime/window): baseline {BL['time_availability_percent']}%, "
    f"fault-tolerant {FT['time_availability_percent']}%",
    f"MTTR = downtime / failure episodes: baseline {BL['MTTR_s']} s, fault-tolerant {FT['MTTR_s']} s "
    f"(episodes: {BL['failure_episodes']} and {FT['failure_episodes']})",
    f"MTTF = uptime / failure episodes: baseline {BL['MTTF_s']} s, fault-tolerant {FT['MTTF_s']} s",
    f"MTBF = MTTF + MTTR: baseline {BL['MTBF_s']} s, fault-tolerant {FT['MTBF_s']} s",
    f"Observed failure rate: baseline {BL['failure_rate_per_hour']} per hour, "
    f"fault-tolerant {FT['failure_rate_per_hour']} per hour",
])
p("Request availability improved from 96.06% to 99.13% and estimated time availability more than "
  "doubled. The absolute availability values look extreme because the observation window is short "
  "and packed with deliberately injected failures. The honest comparison is relative: under the "
  "same injections the fault-tolerant system lost 62 requests out of 7113 while the baseline lost "
  "137 out of 3481, and it recovered about twice as fast (9.5 s versus 19.2 s MTTR).")
h2("4.2 Per experiment results")
for r in CMP:
    lines([f"{r['experiment']}: baseline {r['baseline_success_rate']}% "
           f"({r['baseline_failed']} failed of {r['baseline_requests']}), "
           f"fault-tolerant {r['ft_success_rate']}% "
           f"({r['ft_failed']} failed of {r['ft_requests']})"])
p("In E1, E2 and E4 the baseline dropped to 0%: with one route and one database any component "
  "failure stopped the whole service. The fault-tolerant system served 100% in E1 and E4 and "
  "99.37% in E2, where the six failed requests were the window while the standby was being "
  "promoted. E3 is 100% in both configurations because chaos injection is disabled in baseline "
  "mode, so this experiment shows the timeout and retry behaviour rather than a before/after "
  "difference. E5 and E6 show the remaining gains: interrupted payments improved from 92.8% to "
  "99.5%, high load from 97.3% to 99.0%.")
h2("4.3 Reliability model")
p(f"A simple series-parallel model was built from the measurements. With single instance "
  f"reliabilities R(app) = {RBD['R_app_instance']} and R(db) = {RBD['R_db_node']} observed during "
  f"the campaign, two replicas give R(app-ft) = 1 - (1-{RBD['R_app_instance']})^2 = "
  f"{RBD['R_app_subsystem_ft']} and R(db-ft) = 1 - (1-{RBD['R_db_node']})^2 = "
  f"{RBD['R_db_subsystem_ft']}. The chain gateway x app-ft x db-ft x network gives "
  f"{RBD['R_system_baseline']} for the baseline topology and {RBD['R_system_ft']} for the "
  f"replicated topology, an improvement of {RBD['improvement_percent']} percentage points. The "
  f"measured request availability (99.13%) is close to the model prediction (98.91%), which is a "
  f"reasonable match for a system of this size.")

# ---------------------------------------------------------------- 5
h1("5. System Architecture")
p("The system is a set of Spring Boot (Java 21) services built with Maven and deployed with "
  "Docker Compose. Each service runs as two containers and the gateway routes to both replicas "
  "round robin. PostgreSQL runs twice: a primary container that accepts writes and a standby "
  "container that mirrors them. Communication is HTTP/JSON.")
p("The gateway forwards /api/students, /api/payments, /api/records and /api/timetable to the "
  "service replicas. Every service exposes /actuator/health for Docker health checks and a "
  "/chaos endpoint for failure injection. Payment state lives in four tables: accounts, payments, "
  "payment_journal (append-only state transitions) and payment_checkpoints (periodic committed "
  "totals).")

# ---------------------------------------------------------------- 6
h1("6. Hardware Fault-Tolerance Design")
lines([
    "- Service replication: 2 containers per service behind the gateway. Covers instance crash "
    "and paused node (E1, E4).",
    "- Database replication: application-level dual write to primary and standby. Covers standby "
    "loss (transparent) and primary loss (promotion).",
    "- Automatic failover: after the primary stops answering, the payment service promotes the "
    "standby and continues writing there (E2).",
    "- Health checks: Docker healthcheck on /actuator/health for every container.",
    "- Restart policy: restart unless-stopped brings killed containers back automatically.",
    "- Load balancing: round robin routing with skipping of dead replicas.",
])
p("ECC memory and RAID storage are documented as platform-level assumptions: in a real "
  "deployment they would sit under the database nodes. In this project the dual database setup "
  "plays the role of the primary/standby pair and directly demonstrates that failure of one node "
  "does not stop the service.")

# ---------------------------------------------------------------- 7
h1("7. Software Fault-Tolerance Design")
p("Six mechanisms are implemented (the assignment requires at least four):")
numbered([
    "Retry with exponential backoff. Three attempts with 200/400/800 ms backoff in the gateway "
    "forward loop, in the DbRetry helper and in payment processing.",
    "Timeouts. 6 s gateway timeout, 2 s inter-service timeout. A hung replica is abandoned and "
    "the next attempt goes to the other replica.",
    "Circuit breaker (Resilience4j) per downstream service. Opens after 50% failures over 8 "
    "calls, probes again after 10 s, and fails fast with a clear message while open.",
    "Idempotent payment processing. Payments carry a UNIQUE idempotency key. A replayed payment "
    "returns the original COMMITTED status instead of charging again.",
    "Checkpointing and journal recovery. Every payment state transition is appended to "
    "payment_journal. After every commit a checkpoint stores the committed count and amount. The "
    "/payments/recovery endpoint reports interrupted transactions and consistency state.",
    "Graceful degradation. If the record verification dependency fails, the timetable service "
    "still returns the schedule with status DEGRADED.",
])
p("Rollback is covered by explicit transaction management: a payment debits the account and "
  "inserts the payment row in one SQL transaction, so a failure rolls both back atomically. The "
  "journal records PENDING, RETRY_WAIT, COMMITTED, REJECTED or FAILED for every state change.")

# ---------------------------------------------------------------- 8
h1("8. Implementation")
h2("8.1 Project structure")
lines([
    "midterm/",
    "  common/            chaos injection, replicated DB, retry, backoff",
    "  student-service/   registration and lookup (2 replicas)",
    "  payment-service/   payments: idempotency, journal, checkpoints, failover",
    "  records-service/   grades and transcripts with retry",
    "  timetable-service/ schedule with graceful degradation",
    "  gateway/           routing, retry, circuit breaker, stats",
    "  scripts/           workload, experiments, analysis, report generator",
    "  experiments/       results for both modes, analysis.json, request logs",
])
h2("8.2 Key implementation points")
bullets([
    "ChaosController: every service exposes POST /chaos/delay|error|crash with delay and "
    "probability parameters. In baseline mode the endpoints answer ignored, so both "
    "configurations run the same image and only the mode differs.",
    "ReplicatedDatabase: one class owns all database access. Reads prefer the standby, writes go "
    "to the primary and mirror to the standby. If the primary dies, the payment service calls "
    "promoteStandby() which makes the standby the write target after a promotion marker write.",
    "GatewayController.forward(): circuit breaker check, then up to three attempts across "
    "replicas with backoff. Per-service statistics are exported at /api/stats and are the data "
    "source for the availability calculations.",
    "Idempotency: payments are stored with a UNIQUE idempotency key. A replay returns DUPLICATE "
    "with the original result.",
])
h2("8.3 How to run")
lines([
    "mvn package -DskipTests",
    "docker compose build",
    "RESILIENCE_MODE=ft docker compose up -d",
    "python3 scripts/run_experiments.py experiments/ft ft",
    "python3 scripts/analyze.py",
])
p("For the baseline configuration the stack is restarted with RESILIENCE_MODE=baseline, "
  "DB_REPLICATION=false and gateway routes pointing to a single instance of each service.")

# ---------------------------------------------------------------- 9
h1("9. Experimental Methodology")
p("Each experiment follows the same protocol: reset gateway statistics, run a short clean "
  "workload, inject the failure while the workload keeps sending requests, observe the recovery "
  "and stop. The workload generator runs several threads that send requests through the gateway "
  "at full speed, and logs every request with a timestamp, latency, HTTP code and outcome. From "
  "these logs the analysis script computes success rates, failure windows and downtime estimates.")
lines([
    "E1 app crash: docker kill university-ft-student-1-1, load 4 workers for 25 s",
    "E2 DB failure: docker stop university-ft-postgres-primary-1, load 3 workers for 35 s",
    "E3 timeout: POST /chaos/delay?delayMillis=2500 on payment-1, load 3 workers for 20 s",
    "E4 node failure: docker pause university-ft-records-1-1, load 3 workers for 20 s",
    "E5 interrupted payment: docker kill university-ft-payment-1-1 mid transaction, then replay "
    "with the same idempotency key",
    "E6 high load: no failure, 24 concurrent workers for 30 s",
])
p("Consistency checks: after E2 the recovery endpoint verifies that committed counts and amounts "
  "agree between payments and payment_checkpoints. After E5 the replayed payment must return the "
  "original status without a second charge.")

# ---------------------------------------------------------------- 10
h1("10. Results and Comparison")
p("Summary of the fault-tolerant run:")
lines([
    "E1 app crash: 0 of 34 requests failed. The gateway skipped the dead replica, the container "
    "was restarted by Compose. No data loss.",
    "E2 DB failure: 6 of 956 failed. Standby promotion took about 5.3 s per instance, the journal "
    "stayed intact, consistency verified by /payments/recovery.",
    "E3 delay: 0 of 47 failed. The timeout and retry absorbed the delay, average latency rose "
    "from 0.01 s to 1.30 s.",
    "E4 node pause: 0 of 24 failed. Round robin skipped the paused replica automatically.",
    "E5 interrupted payment: 6 of 1229 failed. The replay with the same key returned DUPLICATE, "
    "the original amount stayed intact, no double charge.",
    "E6 high load: 50 of 4823 failed (98.96% success). Errors only at saturation peaks.",
])
p("Baseline behaviour: E1, E2 and E4 ended with 0% success because the only route or the only "
  "database was gone. E5 kept 92.79% because only the in-flight request and retries were lost. "
  "E6 kept 97.30% with a single instance. E3 is not comparable because chaos injection is "
  "disabled in baseline mode.")
p("Recovery evidence: the journal recovery endpoint reported 1226 committed payments, total "
  "2 980 407 KZT, last checkpoint id 1235, recovery time 33 ms. Standby promotion was logged as "
  "FAILOVER primary database unreachable, promoting standby, followed by the payment completing "
  "on the promoted node. The replayed payment returned DUPLICATE with originalStatus COMMITTED, "
  "which directly proves that idempotency prevented a double charge.")

# ---------------------------------------------------------------- 11
h1("11. Discussion and Limitations")
bullets([
    "Application-level dual write is not a real distributed database. During the primary outage "
    "writes hit only the promoted standby, and when the old primary comes back the nodes can "
    "diverge. A production design would use PostgreSQL streaming replication with a failover tool.",
    "Chaos endpoints inject faults at application level. Host kernel faults, Docker daemon faults "
    "and disk corruption are only approximated by container stop and pause.",
    "Time availability is estimated from failed-request windows and is only an estimate; request "
    "availability is exact.",
    "Circuit breaker parameters were chosen for experiment-scale traffic and would need tuning in "
    "production.",
    "E3 is symmetric because chaos injection is deliberately disabled in baseline mode.",
    "With two replicas the system tolerates one failure per service but not two.",
])

# ---------------------------------------------------------------- 12
h1("12. Conclusion")
p("The project met its success criteria with measured evidence. Failures were detected, their "
  "impact was limited (request availability 99.13% versus 96.06% baseline, 100% success in three "
  "of six scenarios under active failure), and affected services recovered automatically (MTTR "
  "9.5 s versus 19.2 s) without data loss or duplication. The interrupted payment was replayed "
  "idempotently and the journal verified consistency in 33 ms. The main lesson confirmed by the "
  "experiments is that redundancy only pays off together with detection and coordination: two "
  "replicas without health checks and failover logic would not have survived E1 and E2, and a "
  "standby database without promotion logic would just be a second copy of the outage.")

# ---------------------------------------------------------------- 13
h1("13. References")
numbered([
    "A. Avizienis, J.-C. Laprie, B. Randell, C. Landwehr, Basic Concepts and Taxonomy of "
    "Dependable and Secure Computing, IEEE TDSC, 2004.",
    "Spring Boot Documentation, spring.io/projects/spring-boot",
    "Resilience4j Documentation, resilience4j.readme.io",
    "Docker Compose Documentation, docs.docker.com/compose/",
    "PostgreSQL 16 Documentation, postgresql.org/docs/16/",
    "M. Kleppmann, Designing Data-Intensive Applications, O'Reilly, 2017.",
    "Course materials: Fault Tolerance and Dependable Computing, AITU, 2026.",
])

# ---------------------------------------------------------------- 14
h1("14. Appendix")
h2("14.1 Payment service log excerpt")
mono_src = ""
for name in ("payment_excerpt.log", "payment_failover_excerpt.log"):
    path = os.path.join(EX, "logs", name)
    if os.path.exists(path):
        mono_src += open(path).read() + "\n"
par = doc.add_paragraph()
r = par.add_run(mono_src.strip())
r.font.name = "Consolas"
r.font.size = Pt(8)
h2("14.2 Gateway log excerpt")
path = os.path.join(EX, "logs", "gateway_excerpt.log")
if os.path.exists(path):
    par = doc.add_paragraph()
    r = par.add_run(open(path).read().strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8)
h2("14.3 Demonstration script")
numbered([
    "docker kill university-ft-student-1-1, then POST /api/students through the gateway. The "
    "request is served by student-2, then docker compose up -d student-1 restores the replica.",
    "docker stop university-ft-postgres-primary-1, then POST /api/payments. The log shows FAILOVER "
    "and the payment commits on the promoted standby in about 5 s. docker start restores the primary.",
    "POST the same payment twice with the same idempotency key. First answer COMMITTED, second "
    "answer DUPLICATE with the original paymentId, which proves there is no double charge.",
])

out = os.path.join(os.path.expanduser("~"), "Desktop", "Fault tolerance",
                   "MIDTERM_FTR_Report_Kaztayev_Rauan.docx")
doc.save(out)
print("saved", out)
